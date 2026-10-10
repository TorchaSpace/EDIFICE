"""Raporlama paketleri (Excel): ESG veri paketi, EKB hazırlık veri sayfası, portföy özeti. UI'dan bağımsız.

Bunlar resmî beyan ya da belge DEĞİLDİR: uzmanın/raporlayıcının işini hızlandıran, yöntemi ve kaynağı açık veri paketleridir."""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .evidence import SOURCES
from .models import UtilityType
from .quality import assess
from .service import Project

MONTHS = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
HEAD = PatternFill("solid", fgColor="0B3D2E")
WARN = Font(italic=True, color="9C3B00")


def _sheet(wb, title, rows, widths=None, header_row=None):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(list(r))
    if header_row:
        for c in ws[header_row]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def build_esg_package(project: Project, path: str) -> str:
    b, k, a = project.building, project.kpis(), project.assumptions
    ef_e, ef_g = a.emission_factor_kg_per_kwh[UtilityType.ELECTRICITY], a.emission_factor_kg_per_kwh[UtilityType.GAS]
    scope1 = k.gas_kwh * ef_g / 1000
    scope2 = k.electricity_kwh * ef_e / 1000
    q = assess(project)
    yoy = project.yoy()
    wb = Workbook()
    wb.remove(wb.active)
    rows = [["ESG veri paketi — " + b.name], [f"Raporlama yılı: {project.year} · Hazırlanma tarihi: {date.today():%d.%m.%Y}"], [],
            ["Gösterge", "Değer", "Birim", "Not"],
            ["Brüt kullanım alanı", b.floor_area_m2, "m²", ""],
            ["Toplam enerji tüketimi", round(k.total_energy_kwh), "kWh", "elektrik + doğalgaz"],
            ["  Elektrik", round(k.electricity_kwh), "kWh", ""],
            ["  Doğalgaz", round(k.gas_kwh), "kWh", ""],
            ["Enerji yoğunluğu (EUI)", round(k.eui_kwh_m2, 1), "kWh/m²·yıl", ""],
            ["Kapsam 1 (doğalgaz, doğrudan)", round(scope1, 2), "tCO₂e", f"faktör {ef_g} kgCO₂e/kWh (IPCC 2006, net kalorifik değer)"],
            ["Kapsam 2 (elektrik, konum bazlı)", round(scope2, 2), "tCO₂e", f"faktör {ef_e} kgCO₂e/kWh (ETKB 2023, dağıtım hattı)"],
            ["Toplam (Kapsam 1+2)", round(scope1 + scope2, 2), "tCO₂e", "Kapsam 3 ve soğutucu akışkan kaçakları dahil değildir"],
            ["Karbon yoğunluğu", round(k.carbon_kg_m2, 1), "kgCO₂e/m²·yıl", ""],
            ["Su tüketimi", round(k.water_m3), "m³", ""],
            ["Su yoğunluğu", round(k.water_m3_m2, 2), "m³/m²·yıl", ""],
            ["Toplam enerji ve su maliyeti", round(k.total_cost), "₺", "faturadan; tutar boşsa varsayılan tarife"],
            ["Tahmini enerji sınıfı", project.rating()["class"], "", "BEP-TR ölçeği, GÖSTERGEDİR; resmî Enerji Kimlik Belgesi değildir"],
            ["Girdi verisi güvenilirliği", f"{q.level} ({q.score:.0f}/100)", "", f"{len(q.issues)} uyarı; ayrıntı «Veri kalitesi» sayfasında"]]
    if yoy:
        rows += [[], ["Önceki yıla göre değişim", "", "", f"{project.previous_year()} → {project.year}"]]
        for key, label in (("energy", "Enerji"), ("carbon", "Karbon"), ("water", "Su"), ("cost", "Maliyet")):
            rows.append([f"  {label}", round(yoy[key], 1), "%", ""])
    _sheet(wb, "Özet", rows, [38, 18, 16, 70], header_row=4)
    wb["Özet"]["A1"].font = Font(bold=True, size=14)
    monthly = [["Yıl", "Ay", "Elektrik (kWh)", "Doğalgaz (kWh)", "Su (m³)", "Tutar (₺)", "Kapsam 2 (tCO₂e)", "Kapsam 1 (tCO₂e)"]]
    for y in sorted({r.year for r in project.readings}):
        e, g, w_ = (project.monthly(y, u) for u in (UtilityType.ELECTRICITY, UtilityType.GAS, UtilityType.WATER))
        cost = project.monthly_cost(y)
        for m in range(12):
            monthly.append([y, MONTHS[m], round(e[m]), round(g[m]), round(w_[m]), round(cost[m]), round(e[m] * ef_e / 1000, 3), round(g[m] * ef_g / 1000, 3)])
    _sheet(wb, "Aylık", monthly, [8, 12, 16, 16, 12, 14, 18, 18], header_row=1)
    method = [["Yöntem ve kaynaklar"], [],
              ["Kapsam 1 = doğalgaz kWh × " + str(ef_g) + " kgCO₂e/kWh; Kapsam 2 (konum bazlı) = elektrik kWh × " + str(ef_e) + " kgCO₂e/kWh."],
              ["Kapsam 3, soğutucu akışkan kaçakları ve yerinde yenilenebilir üretim ayrıca raporlanmamıştır."],
              ["Bu dosya resmî bir beyan (TSRS/CSRD/GRESB) değildir; raporlayıcıya hazır veri ve açık yöntem sağlar."], []]
    for key in ("ETKB_EF2023", "IPCC2006", "ES2024", "CSB_EKB"):
        s = SOURCES.get(key)
        if s:
            method.append([f"{s['cite']}  [{s['level']}]  {s['url']}"])
    _sheet(wb, "Yöntem ve kaynaklar", method, [140])
    wb["Yöntem ve kaynaklar"]["A1"].font = Font(bold=True, size=13)
    dq = [["Önem", "Alan", "Bulgu"]] + [[i.severity, i.area, i.message] for i in q.issues] if q.issues else [["Önem", "Alan", "Bulgu"], ["-", "-", "Sorun bulunmadı"]]
    _sheet(wb, "Veri kalitesi", dq, [12, 14, 120], header_row=1)
    wb.save(path)
    return path


def build_ekb_sheet(project: Project, path: str) -> str:
    b, k = project.building, project.kpis()
    r = project.rating()
    wb = Workbook()
    wb.remove(wb.active)
    rows = [["EKB hazırlık veri sayfası — " + b.name], [], ["Alan", "Değer"],
            ["Bina adı", b.name], ["Adres", b.address], ["Kullanım türü", b.use_type], ["Brüt kullanım alanı (m²)", b.floor_area_m2],
            ["Kat sayısı", b.floors], ["Yapım yılı", b.year_built], ["Kullanıcı sayısı", b.occupants],
            ["Enlem / boylam", f"{b.lat}, {b.lon}" if b.lat is not None else ""], [],
            ["Yıllık tüketim (baz yıl " + str(project.year) + ")", "Değer"],
            ["Elektrik (kWh)", round(k.electricity_kwh)], ["Doğalgaz (kWh)", round(k.gas_kwh)], ["Su (m³)", round(k.water_m3)],
            ["Nihai enerji yoğunluğu (kWh/m²·yıl)", round(k.eui_kwh_m2, 1)], [],
            ["Gösterge sınıfı (yaklaşık)", r["class"]], ["Kıyas EUI (ENERGY STAR medyanı)", round(r["benchmark"], 1)],
            [], ["ÖNEMLİ: Bu sayfa resmî Enerji Kimlik Belgesi DEĞİLDİR. EKB, yetkili uzmanca BEP-TR yazılımıyla; bina kabuğu, tesisat ve "
                 "birincil enerji hesabıyla düzenlenir. Bu sayfa uzmana veri hazırlığı içindir; yaklaşık sınıf yalnız göstergedir."]]
    ws = _sheet(wb, "Bina ve tüketim", rows, [44, 40], header_row=3)
    ws["A13"].font = Font(bold=True)
    ws[f"A{len(rows)}"].font = WARN
    ws.merge_cells(start_row=len(rows), start_column=1, end_row=len(rows), end_column=2)
    ws.row_dimensions[len(rows)].height = 60
    eq = [["Kategori", "Ekipman", "Kurulum yılı", "Durum (1-5)", "Not"]] + [[e.category, e.name, e.year_installed, e.condition, e.notes] for e in project.equipment]
    _sheet(wb, "Ekipman envanteri", eq, [16, 36, 14, 12, 40], header_row=1)
    wb.save(path)
    return path


def build_portfolio_workbook(projects: list[Project], path: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    head = ["Bina", "Kullanım", "Alan (m²)", "EUI (kWh/m²)", "Sınıf (yaklaşık)", "Sağlık skoru", "Karbon (tCO₂e)", "Maliyet (₺)",
            "Uygun öneri CAPEX (₺)", "Yıllık tasarruf (₺)", "NPV (₺)", "Geri ödeme (yıl)", "Veri güvenilirliği"]
    rows = [head]
    tot = [0.0] * 5
    for p in projects:
        k, h = p.kpis(), p.health()
        codes = p.applicable_codes()
        sc = p.scenario(codes) if codes else None
        f = p.finance(codes) if codes else None
        q = assess(p)
        pay = sc.payback_years if sc and sc.payback_years != float("inf") else None
        rows.append([p.building.name, p.building.use_type, p.building.floor_area_m2, round(k.eui_kwh_m2, 1), p.rating()["class"], round(h.total),
                     round(k.carbon_kg / 1000, 1), round(k.total_cost), round(sc.capex) if sc else 0, round(sc.annual_saving) if sc else 0,
                     round(f.npv) if f else 0, round(pay, 1) if pay else None, f"{q.level} ({q.score:.0f})"])
        for i, v in enumerate((p.building.floor_area_m2, k.carbon_kg / 1000, k.total_cost, sc.capex if sc else 0, sc.annual_saving if sc else 0)):
            tot[i] += v
    rows.append(["TOPLAM", "", round(tot[0]), "", "", "", round(tot[1], 1), round(tot[2]), round(tot[3]), round(tot[4]), "", "", ""])
    ws = _sheet(wb, "Portföy", rows, [28, 14, 12, 14, 14, 12, 14, 14, 18, 18, 14, 14, 18], header_row=1)
    for c in ws[len(rows)]:
        c.font = Font(bold=True)
    wb.save(path)
    return path
