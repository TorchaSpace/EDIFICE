"""Excel şablonu üretimi ve doldurulmuş dosyanın okunması (UI'dan bağımsız)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .models import UtilityType
from .validation import (EQUIPMENT_CATEGORIES, MONTH_NAMES, USE_TYPES, ValidationError, build_from_inputs,
                         parse_number)

LOGO = Path(__file__).resolve().parent / "assets" / "logo_crop.png"
SHEET_INFO, SHEET_USAGE, SHEET_EQ, SHEET_HELP = "Bina", "Tüketim", "Ekipman", "Talimatlar"

NAVY, GREEN = "0B2A80", "0DDD96"
REQ = PatternFill("solid", fgColor="FFF3C4")   # zorunlu: sarı
OPT = PatternFill("solid", fgColor="F1F4F3")   # isteğe bağlı: gri
HEAD = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="D5DDDA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_HEAD = Font(name="Arial", bold=True, color="FFFFFF", size=10)
F_BODY = Font(name="Arial", size=10)
F_HINT = Font(name="Arial", size=9, italic=True, color="6B7A75")
F_TITLE = Font(name="Arial", bold=True, size=16, color=NAVY)

# Sabit yerleşim (okuyucu bu hücrelere bakar)
INFO_ROWS = {"name": 3, "use_type": 4, "address": 5, "area": 6, "year_built": 7, "floors": 8, "occupants": 9}
BASE_YEAR_CELL = "B3"
BASE_FIRST_ROW, PREV_FIRST_ROW = 6, 22
EQ_FIRST_ROW, EQ_LAST_ROW = 4, 33
USAGE_HEADERS = ["Ay", "Elektrik (kWh) *", "Elektrik tutarı (₺)", "Doğalgaz (kWh) *", "Doğalgaz tutarı (₺)",
                 "Su (m³) *", "Su tutarı (₺)"]


def _style_header(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD, F_HEAD, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _cell(ws, ref, value=None, fill=None, fmt=None, align="right"):
    c = ws[ref]
    c.value = value
    c.font, c.border = F_BODY, BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def _help_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    try:
        from openpyxl.drawing.image import Image
        img = Image(str(LOGO))
        img.height, img.width = 44, 44 * img.width / img.height
        ws.add_image(img, "B2")
    except Exception:
        ws["B2"] = "EDIFI'CE"
        ws["B2"].font = F_TITLE
    ws["B5"] = "Bina veri şablonu"
    ws["B5"].font = F_TITLE
    ws["B6"] = "Bu dosyayı doldurup EDIFI'CE'de  Bina Ekle ▸ Excel dosyası seç  adımıyla yükleyin."
    ws["B6"].font = Font(name="Arial", size=10, color="6B7A75")
    rows = [
        ("Nasıl doldurulur?", ""),
        ("1 · Bina", "Bina adı, alan, yapım yılı ve kat sayısını yazın. Kullanım tipini listeden seçin."),
        ("2 · Tüketim", "Baz yılı yazın (son tam 12 ay). Faturalardaki aylık kWh ve m³ değerlerini girin. "
                        "Tutar (₺) sütunları boş bırakılabilir; boşsa uygulamadaki varsayılan tarife kullanılır."),
        ("   Önceki yıl", "İsteğe bağlı. Doldurursanız trend (geçen yıla göre değişim) hesaplanır; "
                          "doldurmayacaksanız tamamen boş bırakın."),
        ("3 · Ekipman", "İsteğe bağlı. HVAC, aydınlatma ve bina kabuğu ekipmanlarını satır satır ekleyin. "
                        "Durum: 1 çok kötü, 5 çok iyi. Ekipman girerseniz öneriler binaya göre önceliklenir."),
        ("Renkler", "SARI hücreler zorunlu, GRİ hücreler isteğe bağlıdır."),
        ("Sayı yazımı", "1234,5  ·  1.234,5  ·  1234.5 biçimlerinin hepsi okunur. Birim ve ₺ işareti yazmayın."),
        ("Dikkat", "Sayfa adlarını, başlık satırlarını ve satır/sütun düzenini değiştirmeyin; yeni sütun eklemeyin."),
    ]
    for i, (a, b) in enumerate(rows, start=8):
        ws.cell(row=i, column=2, value=a).font = Font(name="Arial", bold=True, size=10, color=NAVY if i == 8 else "000000")
        c = ws.cell(row=i, column=3, value=b)
        c.font, c.alignment = F_BODY, Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30 if len(b) > 90 else 18
    ws["B8"].font = Font(name="Arial", bold=True, size=12, color=NAVY)


def _info_sheet(ws, data: dict | None):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 56
    ws["A1"] = "Bina bilgileri"
    ws["A1"].font = F_TITLE
    fields = [
        ("Bina adı *", "name", True, "Raporlarda ve bina listesinde görünecek ad"),
        ("Kullanım tipi", "use_type", False, "Listeden seçin"),
        ("Adres", "address", False, "İl, ilçe ve açık adres"),
        ("Brüt kullanım alanı (m²) *", "area", True, "Tüm katların toplam alanı"),
        ("Yapım yılı *", "year_built", True, "Örn. 1998"),
        ("Kat sayısı *", "floors", True, "Zemin ve bodrum dahil"),
        ("Kullanıcı / çalışan sayısı", "occupants", False, "Günlük ortalama kişi sayısı"),
    ]
    for label, key, req, hint in fields:
        r = INFO_ROWS[key]
        _cell(ws, f"A{r}", label, OPT, align="left").font = Font(name="Arial", bold=True, size=10)
        _cell(ws, f"B{r}", (data or {}).get(key), REQ if req else OPT, align="left")
        c = ws[f"C{r}"]
        c.value, c.font = hint, F_HINT
        ws.row_dimensions[r].height = 22
    ws["B6"].number_format = "#,##0"
    dv = DataValidation(type="list", formula1='"' + ",".join(USE_TYPES) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Listeden bir kullanım tipi seçin.", "Geçersiz seçim"
    ws.add_data_validation(dv)
    dv.add(f"B{INFO_ROWS['use_type']}")
    for key, lo, hi in (("area", 1, 5_000_000), ("year_built", 1800, date.today().year), ("floors", 1, 200), ("occupants", 0, 100000)):
        v = DataValidation(type="decimal" if key == "area" else "whole", operator="between", formula1=lo, formula2=hi,
                           allow_blank=True)
        v.error, v.errorTitle = f"{lo} ile {hi} arasında bir sayı girin.", "Geçersiz değer"
        ws.add_data_validation(v)
        v.add(f"B{INFO_ROWS[key]}")


def _usage_block(ws, first_row: int, rows: list | None, required: bool):
    for c, h in enumerate(USAGE_HEADERS, start=1):
        ws.cell(row=first_row - 1, column=c, value=h)
    _style_header(ws, first_row - 1, 7)
    for m in range(12):
        r = first_row + m
        _cell(ws, f"A{r}", MONTH_NAMES[m], OPT, align="left").font = Font(name="Arial", bold=True, size=10)
        for c in range(2, 8):
            req_col = required and c in (2, 4, 6)
            val = rows[m][c - 2] if rows else None
            _cell(ws, f"{chr(64 + c)}{r}", val, REQ if req_col else OPT, "#,##0")
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
    dv.error, dv.errorTitle = "0 veya daha büyük bir sayı girin (birim yazmayın).", "Geçersiz değer"
    ws.add_data_validation(dv)
    dv.add(f"B{first_row}:G{first_row + 11}")


def _usage_sheet(ws, base_year: int | None, base_rows, prev_rows):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 21
    ws["A1"] = "Aylık tüketim"
    ws["A1"].font = F_TITLE
    _cell(ws, "A3", "Baz yıl *", OPT, align="left").font = Font(name="Arial", bold=True, size=10)
    _cell(ws, "B3", base_year, REQ, "0", "center")
    ws["C3"].value, ws["C3"].font = "Son tam 12 ay (örn. 2025)", F_HINT
    ydv = DataValidation(type="whole", operator="between", formula1=1990, formula2=date.today().year, allow_blank=False)
    ydv.error, ydv.errorTitle = "1990 ile bu yıl arasında bir yıl girin.", "Geçersiz yıl"
    ws.add_data_validation(ydv)
    ydv.add("B3")
    ws["A4"] = "Baz yıl tüketimi"
    ws["A4"].font = Font(name="Arial", bold=True, size=11, color=NAVY)
    _usage_block(ws, BASE_FIRST_ROW, base_rows, True)
    _cell(ws, "A19", "Önceki yıl (isteğe bağlı)", OPT, align="left").font = Font(name="Arial", bold=True, size=10)
    _cell(ws, "B19", "=B3-1", OPT, "0", "center")
    ws["C19"].value, ws["C19"].font = "Trend karşılaştırması için; kullanmayacaksanız boş bırakın", F_HINT
    ws["A20"] = "Önceki yıl tüketimi"
    ws["A20"].font = Font(name="Arial", bold=True, size=11, color=NAVY)
    _usage_block(ws, PREV_FIRST_ROW, prev_rows, False)
    ws.freeze_panes = "B6"


def _eq_sheet(ws, rows: list | None):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (18, 38, 16, 18, 42)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Ekipman envanteri (isteğe bağlı)"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Her satıra bir ekipman yazın. Durum: 1 çok kötü … 5 çok iyi. Boş satırlar yok sayılır."
    ws["A2"].font = F_HINT
    for c, h in enumerate(["Kategori", "Ekipman adı", "Kurulum yılı", "Durum (1-5)", "Not"], start=1):
        ws.cell(row=3, column=c, value=h)
    _style_header(ws, 3, 5)
    for i, r in enumerate(range(EQ_FIRST_ROW, EQ_LAST_ROW + 1)):
        row = rows[i] if rows and i < len(rows) else (None,) * 5
        for c, v in enumerate(row, start=1):
            _cell(ws, f"{chr(64 + c)}{r}", v, OPT, "0" if c in (3, 4) else None, "left" if c in (1, 2, 5) else "center")
    rng = f"{EQ_FIRST_ROW}:{EQ_LAST_ROW}"
    lists = [("A", DataValidation(type="list", formula1='"' + ",".join(EQUIPMENT_CATEGORIES) + '"', allow_blank=True)),
             ("D", DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)),
             ("C", DataValidation(type="whole", operator="between", formula1=1900, formula2=date.today().year, allow_blank=True))]
    for col, dv in lists:
        dv.error, dv.errorTitle = "Geçerli bir değer girin.", "Geçersiz değer"
        ws.add_data_validation(dv)
        dv.add(f"{col}{EQ_FIRST_ROW}:{col}{EQ_LAST_ROW}")
    ws.freeze_panes = "A4"


def build_template(path: str, example: bool = False) -> str:
    """Boş şablon ya da (example=True) örnek verili şablon yazar."""
    info = base_year = base_rows = prev_rows = eq_rows = None
    if example:
        from . import mock_data
        b = mock_data.building()
        info = dict(name="Örnek Ofis Binası", use_type=b.use_type, address=b.address, area=b.floor_area_m2,
                    year_built=b.year_built, floors=b.floors, occupants=b.occupants)
        readings = mock_data.utility_readings()
        base_year = 2025

        def block(year):
            rows = []
            for m in range(1, 13):
                row = []
                for u in (UtilityType.ELECTRICITY, UtilityType.GAS, UtilityType.WATER):
                    r = next(x for x in readings if x.year == year and x.month == m and x.utility == u)
                    row += [r.consumption, r.cost]
                rows.append(row)
            return rows
        base_rows, prev_rows = block(2025), block(2024)
        eq_rows = [(e.category, e.name, e.year_installed, e.condition, e.notes) for e in mock_data.equipment()]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_HELP
    _help_sheet(ws)
    _info_sheet(wb.create_sheet(SHEET_INFO), info)
    _usage_sheet(wb.create_sheet(SHEET_USAGE), base_year, base_rows, prev_rows)
    _eq_sheet(wb.create_sheet(SHEET_EQ), eq_rows)
    wb[SHEET_HELP].sheet_properties.tabColor = NAVY
    wb[SHEET_INFO].sheet_properties.tabColor = GREEN
    wb[SHEET_USAGE].sheet_properties.tabColor = GREEN
    wb[SHEET_EQ].sheet_properties.tabColor = "9AA8A3"
    wb.save(path)
    return path


# ---------------------------------------------------------------- okuma

def _txt(v) -> str:
    """Hücre değerini parse_number'ın anlayacağı metne çevirir (ondalık ayracı virgül)."""
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v).replace(".", ",")
    return str(v).strip()


def read_workbook(path: str, tariffs: dict | None = None):
    """Doldurulmuş şablonu okur -> (Building, readings, equipment, base_year). Hata varsa ValidationError."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        raise ValidationError(["Dosya okunamadı. EDIFI'CE şablonundan kaydedilmiş geçerli bir .xlsx dosyası seçin."])
    names = {n.lower(): n for n in wb.sheetnames}
    missing = [s for s in (SHEET_INFO, SHEET_USAGE) if s.lower() not in names]
    if missing:
        raise ValidationError([f"Şablonda «{m}» sayfası bulunamadı. Sayfa adlarını değiştirmeyin; şablonu yeniden indirin."
                               for m in missing])
    info_ws, use_ws = wb[names[SHEET_INFO.lower()]], wb[names[SHEET_USAGE.lower()]]
    errors: list[str] = []

    def num(ws, ref, label, integer=False):
        raw = _txt(ws[ref].value)
        if not raw:
            return None
        try:
            v = parse_number(raw)
        except ValueError:
            errors.append(f"{ws.title} sayfası {ref}: {label} geçerli bir sayı değil (“{raw}”)")
            return None
        return int(v) if integer and v is not None else v

    r = INFO_ROWS
    info = dict(name=_txt(info_ws[f"B{r['name']}"].value), address=_txt(info_ws[f"B{r['address']}"].value),
                use_type=_txt(info_ws[f"B{r['use_type']}"].value) or USE_TYPES[0],
                floor_area_m2=num(info_ws, f"B{r['area']}", "Brüt kullanım alanı") or 0,
                year_built=num(info_ws, f"B{r['year_built']}", "Yapım yılı", True) or 0,
                floors=num(info_ws, f"B{r['floors']}", "Kat sayısı", True) or 0,
                occupants=num(info_ws, f"B{r['occupants']}", "Kullanıcı sayısı", True) or 0)
    base_year = num(use_ws, BASE_YEAR_CELL, "Baz yıl", True)
    if not base_year:
        errors.append(f"{use_ws.title} sayfası {BASE_YEAR_CELL}: Baz yıl yazılmalı (örn. {date.today().year - 1})")
        base_year = date.today().year - 1

    def grid(first):
        return [[_txt(use_ws.cell(row=first + m, column=c).value) for c in range(2, 8)] for m in range(12)]

    grids = {base_year: grid(BASE_FIRST_ROW), base_year - 1: grid(PREV_FIRST_ROW)}
    eq_rows = []
    if SHEET_EQ.lower() in names:
        ws = wb[names[SHEET_EQ.lower()]]
        for row in range(EQ_FIRST_ROW, EQ_LAST_ROW + 1):
            cat, name, yr, cond, note = (ws.cell(row=row, column=c).value for c in range(1, 6))
            if not _txt(name):
                eq_rows.append({})
                continue
            try:
                yv = int(parse_number(_txt(yr))) if _txt(yr) else 0
                cv = int(parse_number(_txt(cond))) if _txt(cond) else 3
            except ValueError:
                errors.append(f"{ws.title} sayfası satır {row}: kurulum yılı veya durum sayı değil")
                eq_rows.append({})
                continue
            if cv not in range(1, 6):
                errors.append(f"{ws.title} sayfası satır {row}: Durum 1 ile 5 arasında olmalı")
                cv = 3
            eq_rows.append(dict(category=_txt(cat) if _txt(cat) in EQUIPMENT_CATEGORIES else EQUIPMENT_CATEGORIES[0],
                                name=_txt(name), year_installed=yv, condition=cv, notes=_txt(note)))
    try:
        building, readings, equipment = build_from_inputs(info, grids, eq_rows, base_year, tariffs)
    except ValidationError as e:
        errors.extend(e.errors)
    if errors:
        raise ValidationError(errors)
    return building, readings, equipment, base_year
