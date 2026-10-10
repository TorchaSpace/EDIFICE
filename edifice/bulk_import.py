"""Toplu tüketim içe aktarma: CSV/Excel tablosundan (ay ay fatura dökümü) tüketim kayıtları okur. UI'dan bağımsız.

Desteklenen biçimler:
  geniş : Yıl, Ay, Elektrik kWh, Elektrik ₺, Doğalgaz kWh, Doğalgaz ₺, Su m³, Su ₺   (sütun adları esnek; "Dönem" tek sütun da olur)
  uzun  : Tarih/Dönem, Tür (elektrik/doğalgaz/su), Miktar, Tutar
Ondalık/binlik ayırıcı (1.234,56 ya da 1,234.56), ; , veya sekme ayırıcı, UTF-8 ve Windows-1254 kodlaması otomatik anlaşılır."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from .models import UtilityReading, UtilityType

MONTHS = {"ocak": 1, "subat": 2, "mart": 3, "nisan": 4, "mayis": 5, "haziran": 6, "temmuz": 7, "agustos": 8, "eylul": 9, "ekim": 10,
          "kasim": 11, "aralik": 12, "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10,
          "nov": 11, "dec": 12, "oca": 1, "sub": 2, "nis": 4, "haz": 6, "tem": 7, "agu": 8, "eyl": 9, "eki": 10, "kas": 11, "ara": 12}
FOLD = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
UTIL_WORDS = {UtilityType.ELECTRICITY: ("elektrik", "electric", "elk"), UtilityType.GAS: ("dogalgaz", "gaz", "gas"), UtilityType.WATER: ("su", "water")}
COST_WORDS = ("tl", "tutar", "bedel", "fatura", "cost", "₺", "try", "lira")


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("İ", "i").replace("I", "ı").lower().translate(FOLD)).strip()


@dataclass
class ImportResult:
    readings: list[UtilityReading] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    rows_total: int = 0
    rows_used: int = 0
    cost_estimated: int = 0

    @property
    def period(self) -> tuple[tuple[int, int], tuple[int, int]] | None:
        ks = sorted({(r.year, r.month) for r in self.readings})
        return (ks[0], ks[-1]) if ks else None


def parse_number(v) -> float | None:
    """'1.234,56', '1,234.56', '1234,5', 1234.5 -> float; boş/geçersiz -> None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v).strip())
    if not s or s in "-.,":
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif "." in s and re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")                      # Türkçe binlik: 1.234 -> 1234
    try:
        return float(s)
    except ValueError:
        return None


def parse_period(v, year_hint: int | None = None) -> tuple[int, int] | None:
    """Tek hücreden (yıl, ay): 2025-03, 03/2025, Mart 2025, 01.03.2025, tarih nesnesi."""
    if isinstance(v, (datetime, date)):
        return v.year, v.month
    s = norm(v)
    if not s:
        return None
    m = re.search(r"(20\d\d|19\d\d)[-/.](\d{1,2})(?!\d)", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(1)), int(m.group(2))
    m = re.search(r"(?<!\d)(\d{1,2})[-/.](20\d\d|19\d\d)", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(2)), int(m.group(1))
    m = re.search(r"(?<!\d)(\d{1,2})[-/.](\d{1,2})[-/.](20\d\d|19\d\d)", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(3)), int(m.group(2))
    y = re.search(r"(20\d\d|19\d\d)", s)
    for word in s.replace(".", " ").split():
        for name, mm in MONTHS.items():
            if word.startswith(name[:3]) and (len(name) <= 3 or word[:4] == name[:4]):
                year = int(y.group(1)) if y else year_hint
                return (year, mm) if year else None
    if s.isdigit() and 1 <= int(s) <= 12 and year_hint:
        return year_hint, int(s)
    return None


def _read_table(path: str) -> list[list]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1254"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1")
    sample = text[:4000]
    delim = max(";,\t", key=lambda d: sample.count(d))
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _classify_header(h) -> tuple[str, UtilityType | None]:
    """Başlık -> (rol, tür): rol = year | month | period | util | amount | cost | consumption | cost_col"""
    n = norm(h)
    if not n:
        return "", None
    util = next((u for u, words in UTIL_WORDS.items() if any(re.search(rf"(?<![a-z]){w}(?![a-z])", n) for w in words)), None)
    is_cost = any(w in n for w in COST_WORDS)
    if util is not None:
        return ("cost" if is_cost else "consumption"), util
    if n in ("yil", "year", "yıl"):
        return "year", None
    if n in ("ay", "month", "ay no"):
        return "month", None
    if any(w in n for w in ("donem", "tarih", "date", "period")):
        return "period", None
    if any(w in n for w in ("tur", "hizmet", "utility", "cins", "kalem")):
        return "util", None
    if is_cost:
        return "cost", None
    if any(w in n for w in ("miktar", "tuketim", "consumption", "kwh", "m3")):
        return "amount", None
    return "", None


def _util_from_text(v) -> UtilityType | None:
    n = norm(v)
    for u, words in UTIL_WORDS.items():
        if any(re.search(rf"(?<![a-z]){w}(?![a-z])", n) for w in words):
            return u
    return None


def parse_file(path: str, tariffs: dict | None = None) -> ImportResult:
    res = ImportResult()
    try:
        table = _read_table(path)
    except Exception as e:
        res.warnings.append(f"Dosya okunamadı: {e}")
        return res
    # başlık satırı: en çok tanınan sütunu olan ilk 10 satır
    best, hdr_i = 0, None
    for i, row in enumerate(table[:10]):
        score = sum(1 for c in row if _classify_header(c)[0])
        if score > best:
            best, hdr_i = score, i
    if hdr_i is None or best < 2:
        res.warnings.append("Başlık satırı tanınamadı. İlk satırda Yıl, Ay, Elektrik kWh, Elektrik ₺, Doğalgaz kWh … gibi sütun adları olmalı.")
        return res
    cols = [_classify_header(c) for c in table[hdr_i]]
    idx = {}
    for j, (role, util) in enumerate(cols):
        if role in ("year", "month", "period", "util", "amount") and role not in idx:
            idx[role] = j
        elif role in ("consumption", "cost") and util is not None:
            idx[(role, util)] = j
        elif role == "cost" and "cost" not in idx:
            idx["cost"] = j
    wide = any(isinstance(k, tuple) and k[0] == "consumption" for k in idx)
    long = "util" in idx and "amount" in idx
    if not (wide or long):
        res.warnings.append("Tüketim sütunları bulunamadı (ör. «Elektrik kWh», «Doğalgaz kWh», «Su m³» ya da Tür + Miktar).")
        return res
    if not ("period" in idx or "month" in idx):
        res.warnings.append("Dönem sütunu bulunamadı (Ay ya da Dönem/Tarih).")
        return res
    tariffs = tariffs or {}
    seen: dict[tuple, int] = {}
    for r_i, row in enumerate(table[hdr_i + 1:], start=hdr_i + 2):
        if not any(str(c).strip() for c in row if c is not None):
            continue
        res.rows_total += 1
        get = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None
        year_hint = None
        if "year" in idx:
            y = parse_number(get("year"))
            year_hint = int(y) if y else None
        per = parse_period(get("period"), year_hint) if "period" in idx else None
        if per is None and "month" in idx:
            mv = get("month")
            per = parse_period(mv, year_hint) if not isinstance(mv, (int, float)) else ((year_hint, int(mv)) if year_hint and 1 <= int(mv) <= 12 else None)
        if per is None:
            res.warnings.append(f"Satır {r_i}: dönem anlaşılamadı, atlandı.")
            continue
        used = False
        entries: list[tuple[UtilityType, float | None, float | None]] = []
        if wide:
            for u in UtilityType:
                if ("consumption", u) in idx:
                    entries.append((u, parse_number(get(("consumption", u))), parse_number(get(("cost", u)))))
        if long:
            u = _util_from_text(get("util"))
            if u is None:
                res.warnings.append(f"Satır {r_i}: tür anlaşılamadı, atlandı.")
            else:
                entries.append((u, parse_number(get("amount")), parse_number(get("cost"))))
        for u, amount, cost in entries:
            if amount is None:
                continue
            if amount < 0:
                res.warnings.append(f"Satır {r_i}: negatif {u.value} değeri atlandı.")
                continue
            if cost is None or cost < 0:
                cost = amount * tariffs.get(u, 0.0)
                res.cost_estimated += 1
            key = (u, per[0], per[1])
            if key in seen:
                res.warnings.append(f"Satır {r_i}: {per[1]}/{per[0]} {u.value} tekrar ediyor; son değer kullanıldı.")
                res.readings = [x for x in res.readings if (x.utility, x.year, x.month) != key]
            seen[key] = r_i
            res.readings.append(UtilityReading(u, per[0], per[1], float(amount), float(cost)))
            used = True
        res.rows_used += used
    if res.cost_estimated:
        res.warnings.append(f"{res.cost_estimated} kayıtta tutar boştu; varsayılan tarifeyle tahmin edildi (maliyet/NPV tahmin olur).")
    return res


def merge_readings(existing: list[UtilityReading], new: list[UtilityReading]) -> tuple[list[UtilityReading], int, int]:
    """Aynı (tür, yıl, ay) kaydı yenisiyle değiştirilir. -> (birleşik liste, eklenen, değiştirilen)"""
    keys = {(r.utility, r.year, r.month): r for r in existing}
    added = replaced = 0
    for r in new:
        k = (r.utility, r.year, r.month)
        replaced += k in keys
        added += k not in keys
        keys[k] = r
    return sorted(keys.values(), key=lambda r: (r.year, r.month, r.utility.value)), added, replaced


def sample_csv(path: str, readings: list[UtilityReading]) -> str:
    """Mevcut binanın son yılından örnek CSV yazar (kullanıcı biçimi görsün diye)."""
    year = max((r.year for r in readings), default=date.today().year - 1)
    by = {(r.utility, r.month): r for r in readings if r.year == year}
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Yıl", "Ay", "Elektrik kWh", "Elektrik ₺", "Doğalgaz kWh", "Doğalgaz ₺", "Su m³", "Su ₺"])
        for m in range(1, 13):
            row = [year, m]
            for u in (UtilityType.ELECTRICITY, UtilityType.GAS, UtilityType.WATER):
                r = by.get((u, m))
                row += [f"{r.consumption:.0f}".replace(".", ","), f"{r.cost:.0f}"] if r else ["", ""]
            w.writerow(row)
    return path
