"""CRREM yol dosyasını (kullanıcının crrem.org'dan kendi indirdiği resmî Excel) okur. UI'dan bağımsız.

ÖNEMLİ — LİSANS: CRREM verisi (CRREM Foundation) dahili karar araçlarında, kaynak gösterilerek serbest kullanılabilir; ancak ticari yazılıma gömmek
ve yeniden dağıtmak ayrı bir License Partner anlaşması ister (https://crrem.org/library/use-of-data/). Bu yüzden uygulama CRREM verisini İÇERMEZ;
kullanıcı dosyayı kendisi indirir (kayıt formu ve kullanım koşulları onunla) ve buraya içe aktarır. Veri yalnızca kullanıcının yerel veritabanında durur.
Türkiye (TR) CRREM kapsamında YOKTUR (V2.01'de 65 ülke/şehir kodu var, TR yok); vekil ülke seçmek kullanıcının kararıdır."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl import load_workbook


class CrremError(Exception):
    pass


@dataclass
class CrremData:
    version: str = ""
    source_file: str = ""
    pathways: dict[str, dict[tuple[str, str], dict[int, float]]] = field(default_factory=lambda: {"ghg": {}, "kwh": {}})
    grid: dict[str, dict[int, float]] = field(default_factory=dict)

    @property
    def countries(self) -> list[str]:
        return sorted({c for d in self.pathways.values() for c, _ in d})

    @property
    def types(self) -> list[str]:
        return sorted({t for d in self.pathways.values() for _, t in d})


def _header_row(rows):
    for i, r in enumerate(rows[:12]):
        if r and r[0] == "Year":
            return i
    return None


def parse(path: str) -> CrremData:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise CrremError(f"Dosya açılamadı: {e}") from e
    data = CrremData(source_file=path.rsplit("/", 1)[-1])
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            for c in r:
                if isinstance(c, str) and re.search(r"(?i)version\s*:", c):
                    data.version = c.strip()
        name = ws.title.lower()
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if "grid" in name:
            _parse_grid(rows, data)
            continue
        metric = "ghg" if "ghg" in name else "kwh" if "kwh" in name else None
        if metric is None:
            continue
        hi = _header_row(rows)
        if hi is None:
            continue
        header = rows[hi]
        for r in rows[hi + 1:]:
            if not r or not isinstance(r[0], (int, float)) or not 1990 < r[0] < 2100:
                continue
            year = int(r[0])
            for j, code in enumerate(header):
                if j == 0 or not isinstance(code, str) or code.count(".") < 2:
                    continue
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    country, ptype = code.split(".")[0], code.split(".")[1]
                    data.pathways[metric].setdefault((country, ptype), {})[year] = float(v)
    if not data.pathways["ghg"] and not data.pathways["kwh"]:
        raise CrremError("Bu dosyada tanınan CRREM yol tablosu bulunamadı (beklenen: «GHGe» ve «kWh» sayfaları, ilk sütun «Year», "
                         "sütun adları ÜLKE.TÜR.ölçüt). Dosya crrem.org'un yeni bir biçimindeyse içe aktarma güncellenmelidir.")
    return data


def _parse_grid(rows, data: CrremData):
    header = None
    for r in rows:
        if not r or r[0] is None:
            continue
        if isinstance(r[0], str) and r[0].lower().startswith("grid"):
            header = [str(c).strip() if c is not None else None for c in r]
            continue
        if header and isinstance(r[0], (int, float)) and 1990 < r[0] < 2100:
            for j, code in enumerate(header):
                if j and code and j < len(r) and isinstance(r[j], (int, float)):
                    data.grid.setdefault(code, {})[int(r[0])] = float(r[j])
