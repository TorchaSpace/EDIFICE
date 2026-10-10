import pytest
from openpyxl import load_workbook

from edifice.models import UtilityType
from edifice.reports_xlsx import build_ekb_sheet, build_esg_package, build_portfolio_workbook
from edifice.service import Project


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_esg_package_numbers_match_engine(tmp_path):
    p = Project.mock()
    out = build_esg_package(p, str(tmp_path / "esg.xlsx"))
    wb = load_workbook(out)
    assert {"Özet", "Aylık", "Yöntem ve kaynaklar", "Veri kalitesi"} <= set(wb.sheetnames)
    rows = {r[0]: r for r in _rows(wb["Özet"]) if r and r[0]}
    k = p.kpis()
    ef = p.assumptions.emission_factor_kg_per_kwh
    assert rows["Kapsam 1 (doğalgaz, doğrudan)"][1] == pytest.approx(k.gas_kwh * ef[UtilityType.GAS] / 1000, abs=0.01)
    assert rows["Kapsam 2 (elektrik, konum bazlı)"][1] == pytest.approx(k.electricity_kwh * ef[UtilityType.ELECTRICITY] / 1000, abs=0.01)
    assert rows["Toplam (Kapsam 1+2)"][1] == pytest.approx(k.carbon_kg / 1000, abs=0.02)
    assert rows["Toplam enerji tüketimi"][1] == round(k.total_energy_kwh)
    monthly = _rows(wb["Aylık"])
    assert len(monthly) == 1 + 12 * len({r.year for r in p.readings})
    method = " ".join(str(r[0]) for r in _rows(wb["Yöntem ve kaynaklar"]) if r[0])
    assert "resmî bir beyan" in method and "ETKB" in method


def test_ekb_sheet_has_inventory_and_disclaimer(tmp_path):
    p = Project.mock()
    wb = load_workbook(build_ekb_sheet(p, str(tmp_path / "ekb.xlsx")))
    text = " ".join(str(c) for r in _rows(wb["Bina ve tüketim"]) for c in r if c)
    assert "DEĞİLDİR" in text and p.building.name in text
    inv = _rows(wb["Ekipman envanteri"])
    assert len(inv) == 1 + len(p.equipment) and inv[1][1] == p.equipment[0].name


def test_portfolio_workbook_totals(tmp_path):
    p1, p2 = Project.mock(), Project.mock()
    p2.building.name, p2.building.floor_area_m2 = "İkinci", 3000
    wb = load_workbook(build_portfolio_workbook([p1, p2], str(tmp_path / "p.xlsx")))
    rows = _rows(wb["Portföy"])
    assert len(rows) == 4 and rows[-1][0] == "TOPLAM"
    assert rows[-1][2] == round(p1.building.floor_area_m2 + 3000)
    assert rows[1][0] == p1.building.name and rows[1][4] == p1.rating()["class"]
