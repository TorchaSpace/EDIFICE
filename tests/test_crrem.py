import pytest
from openpyxl import Workbook

from edifice.crrem import CrremError, parse
from edifice.db import Store
from edifice.engine.stranding import analyze


def synthetic_workbook(path, with_tr=False):
    """CRREM V2.0x dosyasının yapısını taklit eder (sayfa adları, 'Year' başlığı, ÜLKE.TÜR.ölçüt sütunları). Sayılar UYDURMADIR: yalnız ayrıştırıcıyı sınar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A4"] = "Version: v2.01 - 11.01.2023"
    for name, metric, base in (("2 - 1.5 kWh", "kWh-Int", 150.0), ("3 - 1.5C GHGe", "GHG-Int", 60.0)):
        s = wb.create_sheet(name)
        codes = ["ES.OFF." + metric, "DE.OFF." + metric] + (["TR.OFF." + metric] if with_tr else [])
        s.append(["Year"] + codes)
        s.append([None] * (len(codes) + 1))
        for i, y in enumerate(range(2020, 2051)):
            s.append([y] + [base * (1 - 0.03 * i) * (1 + 0.1 * k) for k in range(len(codes))])
    g = wb.create_sheet("Grid EF")
    g.append(["All emission factors excluding transmission & distribution losses"])
    g.append(["Grid EU", "ES", "DE"])
    for i, y in enumerate(range(2020, 2051)):
        g.append([y, 0.25 * (1 - 0.03 * i), 0.4 * (1 - 0.04 * i)])
    wb.save(path)
    return path


def test_parse_structure_and_missing_turkey(tmp_path):
    d = parse(synthetic_workbook(str(tmp_path / "c.xlsx")))
    assert d.version.startswith("Version: v2.01") and d.countries == ["DE", "ES"] and d.types == ["OFF"]
    assert d.pathways["ghg"][("ES", "OFF")][2020] == pytest.approx(60.0) and len(d.pathways["ghg"][("ES", "OFF")]) == 31
    assert d.grid["DE"][2020] == pytest.approx(0.4) and "TR" not in d.countries


def test_unrecognized_file_is_explained(tmp_path):
    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(CrremError) as e:
        parse(str(tmp_path / "x.xlsx"))
    assert "tanınan CRREM yol tablosu bulunamadı" in str(e.value)
    with pytest.raises(CrremError):
        parse(str(tmp_path / "yok.xlsx"))


def test_store_import_roundtrip_and_clear(tmp_path):
    st = Store(":memory:")
    d = parse(synthetic_workbook(str(tmp_path / "c.xlsx")))
    n = st.import_crrem(d)
    assert n == 2 * 2 * 31 and st.crrem_options() == (["DE", "ES"], ["OFF"])
    assert st.crrem_pathway("ghg", "ES", "OFF")[2030] == pytest.approx(60 * (1 - 0.03 * 10))
    assert st.crrem_meta()["file"] == "c.xlsx" and st.crrem_grid("ES")[2020] == pytest.approx(0.25)
    st.import_crrem(d)                                       # tekrar içe aktarma çoğaltmaz
    assert len(st.crrem_pathway("kwh", "DE", "OFF")) == 31
    st.clear_crrem()
    assert st.crrem_options() == ([], []) and st.crrem_meta() is None


PATH = {y: 100 - 3 * (y - 2025) for y in range(2025, 2051)}      # yılda 3 birim azalan yol


def test_stranding_year_without_plan_and_with_plan():
    s = analyze(PATH, 2025, current=100)
    assert s.stranded_year == 2026 and s.excess_total > 0
    # her yıl 4 birim tasarruf (yoldan hızlı): hiç aşmaz
    fast = analyze(PATH, 2025, 100, {y: 4.0 for y in range(2026, 2051)})
    assert fast.stranded_year is None and fast.excess_total == 0
    # tek seferlik 20 birim tasarruf 2027'de: 2026 aşılır, sonra yol düşerek yeniden aşılır
    one = analyze(PATH, 2025, 100, {2027: 20.0})
    assert one.stranded_year == 2026 and one.projected[-1] == pytest.approx(80.0)
    assert analyze({}, 2025, 100).note and analyze(PATH, 2060, 100).note


def test_grid_scale_only_affects_electric_part():
    scale = {y: 1 - 0.05 * (y - 2025) for y in range(2025, 2051)}
    a = analyze(PATH, 2025, 100, elec_part=50, grid_scale=scale)
    assert a.projected[0] == pytest.approx(100) and a.projected[10] == pytest.approx(50 + 50 * (1 - 0.5))
    b = analyze(PATH, 2025, 100, elec_part=0, grid_scale=scale)
    assert b.projected[10] == pytest.approx(100)


def test_ui_panel_imports_and_reports_stranding(tmp_path, monkeypatch):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
    from edifice.ui.main_window import MainWindow
    QApplication.instance() or QApplication([])
    st = Store(":memory:")
    bid = st.seed_demo()
    w = MainWindow(st.load_project(bid), st)
    esg = w.pages[w.idx["Sürdürülebilirlik"]][1]
    from edifice.ui.crrem_panel import CrremPanel
    panel = esg.widget.findChildren(CrremPanel)[0]
    path = synthetic_workbook(str(tmp_path / "crrem.xlsx"))
    monkeypatch.setattr(QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (path, "")))
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
    panel.import_file()
    assert st.crrem_meta()["file"] == "crrem.xlsx"
    _bid, _r, combos, res = panel.row_widgets[0]
    combos[0].setCurrentIndex(combos[0].findData("ES"))
    combos[1].setCurrentIndex(combos[1].findData("OFF"))
    assert "aşıyor" in res.text() or "aşar" in res.text() or "altında" in res.text()
    assert json_sel(st, bid)["country"] == "ES"


def json_sel(st, bid):
    import json
    return json.loads(st.get_setting(f"crrem_sel_{bid}"))
