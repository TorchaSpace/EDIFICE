import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QDialog, QTableWidgetItem

from edifice.ui.building_dialog import make_item

from edifice.db import Store
from edifice.service import Project
from edifice.ui.building_dialog import BuildingDialog
from edifice.ui.main_window import MainWindow
from edifice.ui.report import build_report


def _app():
    return QApplication.instance() or QApplication([])


def test_window_builds_and_scenario_updates():
    _app()
    p = Project.mock()
    w = MainWindow(p)
    sc = w.sub["Mevcut vs Hedef"]
    sc.checks["LED"].setChecked(True)
    assert sc.selected_codes() == ["LED"]
    assert "EDIFI" in build_report(p, ["LED"])


def test_add_building_dialog_flow():
    _app()
    store = Store(":memory:")
    store.seed_demo()
    w = MainWindow(store.load_project(store.latest_id()), store)
    d = BuildingDialog(w)
    d.save()
    assert d.result_data is None and "Bina adı" in d.error.text()
    d.name.setText("Yeni Bina")
    d.area.setValue(2500)
    for r in range(12):
        for c, v in enumerate(["40000", "", "15000", "", "90", ""]):
            d.grids["base"].setItem(r, c, make_item(v, c))
    d.add_equipment_row()
    d.eq.cellWidget(0, 1).inner.setText("Chiller")
    d.use_type.showPopup()
    d.use_type._popup.card.findChildren(type(d.use_type._popup.card.layout().itemAt(1).widget()))[1].click()
    assert d.use_type.currentIndex() == 1
    d.save()
    assert d.result_data is not None
    building, readings, equipment = d.result_data
    bid = store.save_building(building, readings, equipment)
    w.set_project(store.load_project(bid))
    assert w.project.building.name == "Yeni Bina" and store.count() == 2


def test_edit_building_and_settings_flow():
    _app()
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    d = BuildingDialog(w, project=w.project)
    assert d.name.text() == "Demo Ofis Binası" and d.eq.rowCount() == len(w.project.equipment)
    d.name.setText("Düzenlenmiş")
    d.save()
    assert d.result_data is not None and d.result_data[0].name == "Düzenlenmiş"
    assert len(d.result_data[1]) == len(w.project.readings)
    store.update_building(bid, *d.result_data)
    w.set_project(store.load_project(bid))
    assert w.project.building.name == "Düzenlenmiş"
    sp = w.sub["Ayarlar"]
    sp.b_eui.setValue(400)
    sp.opp_widgets[0][2].setValue(30)
    sp.save()
    assert w.project.assumptions.benchmark_eui_kwh_m2 == 400
    assert w.project.opportunities[0].saving_pct == 0.3


def test_scenario_restored_after_reload():
    _app()
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    w.sub["Mevcut vs Hedef"].checks["LED"].setChecked(True)
    w.set_project(store.load_project(bid))
    assert w.sub["Mevcut vs Hedef"].selected_codes() == ["LED"]


def test_pdf_report_generated(tmp_path):
    _app()
    from edifice.ui.report_pdf import build_pdf
    p = Project.mock()
    out = tmp_path / "r.pdf"
    build_pdf(p, ["LED", "CHILLER"], str(out))
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 5000


def test_excel_import_flow(tmp_path, monkeypatch):
    _app()
    from edifice.excel_io import build_template
    from edifice.ui import main_window as mw
    f = tmp_path / "bina.xlsx"
    build_template(str(f), example=True)
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)

    class FakeChoice:
        Accepted = mw.AddChoiceDialog.Accepted
        choice, path = "excel", str(f)
        def __init__(self, parent=None): pass
        def exec(self): return self.Accepted

    seen = {}

    class FakeDialog(mw.BuildingDialog):
        def exec(self):
            seen["review"] = self.review and self.name.text()
            self.save()
            return QDialog.Accepted

    monkeypatch.setattr(mw, "AddChoiceDialog", FakeChoice)
    monkeypatch.setattr(mw, "BuildingDialog", FakeDialog)
    w.add_building()
    assert seen["review"] == "Örnek Ofis Binası"
    assert w.project.building.name == "Örnek Ofis Binası" and store.count() == 2


def test_budget_slider_selects_package_and_finance_updates():
    _app()
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    sc = w.sub["Mevcut vs Hedef"]
    sc.slider.setValue(0)
    assert sc.selected_codes() == []
    sc.slider.setValue(60)           # 3,0 M ₺
    assert sc.selected_codes() and w.project.finance(sc.selected_codes()).npv > 0
    assert store.load_scenario(bid) == sc.selected_codes()


def test_assistant_chat_answers_locally_and_keeps_history():
    import time
    from PySide6.QtWidgets import QApplication
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow

    store = Store(":memory:")
    store.seed_demo()
    w = MainWindow(store.load_project(store.latest_id()), store)
    chat = w.sub["Sohbet"]
    chat.send("Hangi öneriyle başlamalıyım?")
    for _ in range(300):
        QApplication.processEvents()
        if chat.worker.isFinished():
            break
        time.sleep(0.02)
    QApplication.processEvents()
    role, text, _chart, _meta = w.chat_state.display[-1]
    assert role == "assistant" and "geri ödeme" in text
    w.set_project(store.load_project(store.latest_id()))   # bina değişince geçmiş korunur
    assert len(w.chat_state.display) == 2


def test_assistant_actions_charts_and_training_flow():
    import time
    from PySide6.QtWidgets import QApplication
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow

    store = Store(":memory:")
    store.seed_demo()
    w = MainWindow(store.load_project(store.latest_id()), store)
    chat = w.sub["Sohbet"]

    def ask(q):
        chat.send(q)
        for _ in range(400):
            QApplication.processEvents()
            if chat.worker.isFinished():
                break
            time.sleep(0.02)
        for _ in range(5):
            QApplication.processEvents()
    ask("vfd yi planlandı yap 2027")
    assert store.load_projects(w.project.building_id) == {"VFD": ("Planlandı", 2027)}
    ask("led ve vfd yi senaryoda seç")
    assert w.sub["Mevcut vs Hedef"].selected_codes() == ["LED", "VFD"]
    ask("hangi ay en yüksek")
    assert w.chat_state.display[-1][2]["kind"] == "area"
    ask("tesisin ısı pompası var mı")                       # anlaşılmaz -> eğitim listesine düşer
    unknown = store.list_unknown()
    assert [q for _, q in unknown] == ["tesisin ısı pompası var mı"]
    w.sub["Eğitim"].teach(unknown[0][0], unknown[0][1], "equipment")
    assert store.list_unknown() == [] and store.list_examples()
    ask("tesisin ısı pompası var mı")
    assert "Ekipman" in w.chat_state.display[-1][1]          # öğretilen soru artık anlaşılıyor


def test_feedback_clarification_guard_and_reset():
    import time
    from PySide6.QtWidgets import QApplication
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow

    store = Store(":memory:")
    store.seed_demo()
    w = MainWindow(store.load_project(store.latest_id()), store)
    chat, train = w.sub["Sohbet"], w.sub["Eğitim"]

    def ask(q, **kw):
        chat.send(q, **kw)
        for _ in range(400):
            QApplication.processEvents()
            if chat.worker.isFinished():
                break
            time.sleep(0.02)
        for _ in range(5):
            QApplication.processEvents()
    ask("sağlık skorum kaç")
    meta = w.chat_state.display[-1][3]
    assert meta["intent"] == "health"
    chat._feedback(meta, True)                                   # 👍: soru öğrenilir
    assert ("health", "sağlık skorum kaç") in [(i, t) for _, i, t in store.list_examples()]
    chat._feedback(dict(meta, question="bir şey"), False)       # 👎: Eğitim listesine düşer
    assert "bir şey" in [q for _, q in store.list_unknown()]
    # Eğitim: zararlı örnek reddedilir, sıfırlama çalışır
    uid = store.list_unknown()[0][0]
    train.teach(uid, "skorum neden düştü", "equipment")
    assert "skorum neden düştü" not in [t for _, _, t in store.list_examples()] and train.note
    train.reset()
    assert store.list_examples() == []


def test_import_consumption_merges_into_building(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from edifice.bulk_import import sample_csv
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow

    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    before = len(w.project.readings)
    csv_path = tmp_path / "yeni.csv"
    csv_path.write_text("Yıl;Ay;Elektrik kWh;Elektrik ₺\n2027;1;1000;4200\n2027;2;900;3780\n", encoding="utf-8-sig")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(csv_path), ""))
    monkeypatch.setattr(QMessageBox, "exec", lambda self: QMessageBox.Yes)
    w.import_consumption()
    assert len(w.project.readings) == before + 2
    assert any(r.year == 2027 and r.month == 1 and r.consumption == 1000 for r in w.project.readings)


def test_esg_stranded_table_reflects_plan():
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    page = w.pages[w.idx["Sürdürülebilirlik"]][1]
    texts = [lbl.text() for lbl in page.widget.findChildren(__import__("PySide6.QtWidgets", fromlist=["QLabel"]).QLabel)]
    assert any("Binalar hedef yolunu ne zaman aşar" in t for t in texts)
    from PySide6.QtWidgets import QTableWidget
    tables = page.widget.findChildren(QTableWidget)
    cells = [t.item(r, c).text() for t in tables for r in range(t.rowCount()) for c in range(t.columnCount()) if t.item(r, c)]
    assert any(c in ("Şimdi aşıyor", "Hedefte") or c.endswith("'de aşar") for c in cells)


def test_export_packs_write_files(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QFileDialog
    from PySide6.QtGui import QDesktopServices
    from edifice.db import Store
    from edifice.ui.main_window import MainWindow
    store = Store(":memory:")
    bid = store.seed_demo()
    w = MainWindow(store.load_project(bid), store)
    monkeypatch.setattr(QDesktopServices, "openUrl", staticmethod(lambda *_: True))
    for kind in ("esg", "ekb", "portfoy"):
        out = tmp_path / f"{kind}.xlsx"
        monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, o=out, **k: (str(o), "")))
        w.export_pack(kind)
        assert load_workbook(out).sheetnames
