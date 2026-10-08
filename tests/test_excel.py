import pytest
from openpyxl import load_workbook

from edifice.excel_io import SHEET_USAGE, build_template, read_workbook
from edifice.validation import ValidationError


def test_example_template_roundtrip(tmp_path):
    f = tmp_path / "ornek.xlsx"
    build_template(str(f), example=True)
    b, readings, equipment, base_year = read_workbook(str(f))
    assert b.name == "Örnek Ofis Binası" and b.floor_area_m2 == 6000 and base_year == 2025
    assert len(readings) == 72 and len(equipment) == 6


def test_blank_template_reports_required_fields(tmp_path):
    f = tmp_path / "bos.xlsx"
    build_template(str(f))
    with pytest.raises(ValidationError) as e:
        read_workbook(str(f))
    text = " ".join(e.value.errors)
    assert "Bina adı" in text and "alan" in text and "Baz yıl" in text


def test_text_numbers_and_decimal_floats(tmp_path):
    f = tmp_path / "x.xlsx"
    build_template(str(f), example=True)
    wb = load_workbook(f)
    ws = wb[SHEET_USAGE]
    ws["B6"] = "1.234,5"      # metin, TR biçimi
    ws["D6"] = 1.234          # gerçek ondalık sayı (bin ayracı sanılmamalı)
    ws["F6"] = 85
    wb.save(f)
    _, readings, _, _ = read_workbook(str(f))
    jan = {r.utility.value: r.consumption for r in readings if r.year == 2025 and r.month == 1}
    assert jan["electricity"] == 1234.5 and jan["gas"] == 1.234


def test_bad_cell_and_wrong_file(tmp_path):
    f = tmp_path / "bad.xlsx"
    build_template(str(f), example=True)
    wb = load_workbook(f)
    wb[SHEET_USAGE]["B8"] = "abc"
    wb.save(f)
    with pytest.raises(ValidationError) as e:
        read_workbook(str(f))
    assert "abc" in " ".join(e.value.errors)
    junk = tmp_path / "junk.xlsx"
    junk.write_text("not excel")
    with pytest.raises(ValidationError) as e2:
        read_workbook(str(junk))
    assert "okunamadı" in e2.value.errors[0]
